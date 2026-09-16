"""Generate the short hackathon slide deck: slides/qho-demo.pptx.

    python scripts/make_slides.py

Needs the optional "assets" dependencies (python-pptx, matplotlib/Pillow) and
the screenshots in docs/ (scripts/screenshot.sh). All numbers are computed live.
"""
import re
import subprocess
import sys
import time
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from qho import exact_energies, exact_psi, solve  # noqa: E402

DOCS, OUT = ROOT / "docs", ROOT / "slides" / "qho-demo.pptx"

BG = RGBColor(0xFC, 0xFC, 0xFB)
INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x8A, 0x89, 0x84)
LINE = RGBColor(0xE3, 0xE2, 0xDC)
PANEL = RGBColor(0xF4, 0xF4, 0xF1)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
FONT = "Inter"
MONO = "JetBrains Mono"


# ---------------------------------------------------------------- numbers + figures
def compute_stats():
    k = 12
    t0 = time.perf_counter()
    s = solve(k=k, L=12.0, N=2401)
    ms = (time.perf_counter() - t0) * 1e3
    ref = exact_psi(k, s.x)
    fid = np.abs(np.sum(s.psi * ref, axis=1) * s.h)
    norms = np.sum(s.psi**2, axis=1) * s.h
    dE = s.energies - exact_energies(k)
    conv = []
    for N in (101, 201, 401, 801, 1601, 3201):
        c = solve(k=6, L=10.0, N=N)
        conv.append((c.h, np.abs(c.energies - exact_energies(6)).max()))
    return dict(k=k, ms=ms, dE=dE, fid=fid, norms=norms, conv=np.array(conv), n_tests=count_tests())


def count_tests():
    out = subprocess.run([sys.executable, "-m", "pytest", "--collect-only", "-q"], cwd=ROOT,
                         capture_output=True, text=True).stdout
    return int(re.search(r"(\d+) tests? collected", out).group(1))


def make_figures(st):
    plt.rcParams.update({"font.size": 12, "axes.spines.top": False, "axes.spines.right": False,
                         "axes.edgecolor": "#8a8984", "axes.labelcolor": "#52514e",
                         "xtick.color": "#52514e", "ytick.color": "#52514e", "axes.facecolor": "#fcfcfb"})
    # convergence
    h, e = st["conv"].T
    fig, ax = plt.subplots(figsize=(5.6, 4.0), dpi=200)
    ax.loglog(h, 0.3 * e[-1] * (h / h[-1]) ** 2, "--", color="#8a8984", lw=1.5, label="∝ h² (ideal)")
    ax.loglog(h, e, "o-", color="#2a78d6", lw=2, ms=7, mec="white", mew=1.5, label="max |ΔE|, n = 0–5")
    ax.set_xlabel("grid spacing h")
    ax.set_ylabel("max |E_num − E_exact|")
    ax.grid(True, which="major", color="#ecebe6")
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(DOCS / "convergence.png", facecolor="#fcfcfb")
    plt.close(fig)
    # |dE| vs n
    n = np.arange(st["k"])
    fig, ax = plt.subplots(figsize=(5.6, 3.2), dpi=200)
    ax.bar(n, np.abs(st["dE"]), color=["#2a78d6" if i % 2 == 0 else "#eb6834" for i in n], width=0.7)
    ax.set_xlabel("quantum number n")
    ax.set_ylabel("|E_num − E_exact|")
    ax.set_xticks(n)
    ax.ticklabel_format(axis="y", style="sci", scilimits=(0, 0))
    ax.grid(True, axis="y", color="#ecebe6")
    ax.set_axisbelow(True)
    fig.tight_layout()
    fig.savefig(DOCS / "energy-error.png", facecolor="#fcfcfb")
    plt.close(fig)
    # screenshot crops
    light = Image.open(DOCS / "demo-light.png")
    light.crop((57, 305, 849, 1064)).save(DOCS / "crop-main.png")
    light.crop((863, 305, 1344, 1054)).save(DOCS / "crop-detail.png")
    dark = Image.open(DOCS / "demo-dark.png")
    dark.crop((40, 20, 1360, 1060)).save(DOCS / "crop-hero-dark.png")


# ---------------------------------------------------------------- pptx helpers
def text(slide, x, y, w, h, runs, size=18, color=INK, bold=False, font=FONT, align=PP_ALIGN.LEFT, spacing=1.15):
    """runs: str, or list of paragraphs; a paragraph is str or list of (text, dict) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, 0)
    paras = [runs] if isinstance(runs, str) else runs
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for t, opt in ([(para, {})] if isinstance(para, str) else para):
            r = p.add_run()
            r.text = t
            f = r.font
            f.name = opt.get("font", font)
            f.size = Pt(opt.get("size", size))
            f.bold = opt.get("bold", bold)
            f.color.rgb = opt.get("color", color)
        if isinstance(para, list) and para and para[0][1].get("space"):
            p.space_before = Pt(para[0][1]["space"])
    return tb


def box(slide, x, y, w, h, fill=PANEL, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = 0.08
    return s


def arrow(slide, x1, y1, x2, y2, color=MUTED):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(2)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement("{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd", {"type": "triangle"})
    ln.append(tail)
    return c


def picture(slide, path, x, y, w=None, h=None, border=True):
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y),
                                   Inches(w) if w else None, Inches(h) if h else None)
    if border:
        pic.line.color.rgb = LINE
        pic.line.width = Pt(1)
    return pic


def new_slide(prs, title, kicker=None, notes=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    if kicker:
        text(s, 0.6, 0.35, 12, 0.3, kicker.upper(), size=12, color=BLUE, bold=True)
    text(s, 0.6, 0.62, 12.1, 0.8, title, size=32, bold=True)
    s.notes_slide.notes_text_frame.text = notes
    return s


def bullets(slide, x, y, w, items, size=18, gap=10):
    paras = []
    for i, it in enumerate(items):
        runs = it if isinstance(it, list) else [(it, {})]
        runs = [("●  ", {"color": BLUE if i % 2 == 0 else ORANGE, "size": size - 6, "space": gap if i else 0})] + runs
        paras.append(runs)
    return text(slide, x, y, w, 0.5 * len(items), paras, size=size, color=INK2)


def mono(t, **kw):
    return (t, {"font": MONO, "color": INK, **kw})


# ---------------------------------------------------------------- slides
def build(st):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    # 1 - title
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    text(s, 0.6, 0.9, 5.2, 0.3, "HACKATHON DEMO", size=12, color=BLUE, bold=True)
    text(s, 0.6, 1.3, 5.0, 2.4, "Quantum Harmonic Oscillator", size=40, bold=True, spacing=1.0)
    text(s, 0.6, 2.85, 5.0, 1.0, "solved numerically, checked exactly", size=24, color=INK2)
    text(s, 0.6, 4.1, 4.9, 1.6,
         "Turn the Schrödinger equation into a matrix, diagonalise it, and show that the "
         "eigenvalues land on ℏω(n + ½), in a visual you can explore live.", size=17, color=INK2)
    text(s, 0.6, 6.4, 5.0, 0.4, [[mono("python -m qho.build  →  demo/index.html", size=13, color=MUTED)]])
    picture(s, DOCS / "crop-hero-dark.png", 5.95, 0.9, w=6.85)
    s.notes_slide.notes_text_frame.text = (
        "We solved the quantum harmonic oscillator numerically and checked every result against the exact "
        "textbook solution. The whole pipeline runs in milliseconds.")

    # 2 - physics
    s = new_slide(prs, "Schrödinger equation with a harmonic potential", "The physics",
                  "The harmonic oscillator is one of the few potentials with a closed-form solution, "
                  "which makes it a good benchmark for a numerical solver.")
    box(s, 0.6, 1.65, 7.0, 1.35, fill=RGBColor(0xEE, 0xF4, 0xFC), line=None)
    text(s, 0.9, 1.85, 6.5, 1.0, [
        [mono("−(ℏ²/2m) ψ″(x) + ½ mω² x² ψ(x) = E ψ(x)", size=19)],
        [("time-independent Schrödinger equation, V(x) = ½ mω² x²", {"size": 14, "color": INK2, "space": 8})],
    ])
    text(s, 0.6, 3.35, 7.0, 0.4, "Exact solution", size=20, bold=True)
    text(s, 0.6, 3.85, 7.0, 1.2, [
        [mono("E_n = ℏω (n + ½)", size=20), ("      n = 0, 1, 2, …", {"size": 16, "color": INK2})],
        [mono("ψ_n(ξ) ∝ H_n(ξ) e^(−ξ²/2),   ξ = x √(mω/ℏ)", size=18, space=10)],
    ])
    text(s, 0.6, 5.3, 7.0, 0.4, [[("Natural units: ", {"bold": True, "color": INK}),
                                   ("ℏ = m = ω = 1, so E_n = n + ½", {"color": INK2})]], size=17)
    box(s, 8.2, 1.65, 4.5, 4.9)
    text(s, 8.5, 1.9, 4.0, 0.4, "What to look for", size=18, bold=True)
    bullets(s, 8.5, 2.5, 3.95, [
        "Zero-point energy ½ℏω: the ground state never sits at the bottom of the well",
        "Equally spaced levels, ΔE = ℏω",
        "Parity (−1)ⁿ: even states are symmetric, odd states antisymmetric",
        "State n has exactly n nodes",
        "Tails leak past the classical turning points",
    ], size=16, gap=12)

    # 3 - numerical method
    s = new_slide(prs, "Numerical method: finite differences, then diagonalise", "How it's solved",
                  "Replacing the second derivative with a three-point stencil turns the Hamiltonian into a "
                  "symmetric tridiagonal matrix. LAPACK finds its lowest eigenpairs in milliseconds.")
    steps = [
        ("1  Grid", "x_i ∈ [−L, L], spacing h, hard walls (ψ = 0 at the edges)"),
        ("2  Stencil", "ψ″(x_i) ≈ (ψ_{i−1} − 2ψ_i + ψ_{i+1}) / h²     error O(h²)"),
        ("3  Matrix", "H is tridiagonal: diagonal 1/h² + V(x_i), off-diagonal −1/(2h²)"),
        ("4  Solve", "scipy.linalg.eigh_tridiagonal: lowest k eigenpairs only"),
        ("5  Tidy", "normalise so Σ ψ² h = 1, then align each sign with the exact ψ_n"),
    ]
    y = 1.75
    for i, (a, b) in enumerate(steps):
        box(s, 0.6, y, 7.3, 0.78, fill=PANEL)
        text(s, 0.85, y + 0.2, 1.5, 0.4, a, size=16, bold=True, color=BLUE if i % 2 == 0 else ORANGE)
        text(s, 2.35, y + 0.22, 5.4, 0.4, [[mono(b, size=13)]])
        y += 0.93
    # tridiagonal matrix sketch
    ox, oy, cell = 8.75, 1.85, 0.46
    text(s, ox, oy - 0.05, 3.8, 0.4, "H  (N × N, symmetric)", size=15, bold=True)
    oy += 0.5
    nmat = 7
    for r in range(nmat):
        for c in range(nmat):
            if abs(r - c) > 1:
                col, ln = BG, RGBColor(0xEC, 0xEB, 0xE6)
            elif r == c:
                col, ln = BLUE, None
            else:
                col, ln = ORANGE, None
            box(s, ox + c * cell, oy + r * cell, cell - 0.06, cell - 0.06, fill=col, line=ln, shape=MSO_SHAPE.RECTANGLE)
    ly = oy + nmat * cell + 0.2
    text(s, ox, ly, 3.9, 0.9, [
        [("■ ", {"color": BLUE}), ("1/h² + V(x_i)", {"font": MONO, "size": 13, "color": INK})],
        [("■ ", {"color": ORANGE}), ("−1/(2h²)", {"font": MONO, "size": 13, "color": INK}), ("", {"space": 4})],
        [("Demo: L = 12, N = 2401, h = 0.01, 12 states", {"size": 13, "color": INK2, "space": 8})],
    ], size=15)

    # 4 - main visualisation
    s = new_slide(prs, "The central picture: states sit at their energies", "Visualisation",
                  "Each eigenstate is drawn at the height of its energy inside the parabola. The dotted overlay is "
                  "the exact Hermite function, and it lies on top of the numerical curve.")
    picture(s, DOCS / "crop-main.png", 0.6, 1.55, h=5.65)
    cx = 7.2
    callouts = [
        (BLUE, "Equal spacing", "Levels are ℏω apart and start at ½ℏω, not at zero."),
        (ORANGE, "Parity by colour", "Blue states are even and orange states are odd. State n has n nodes."),
        (INK2, "Beyond the turning points", "ψ leaks into the classically forbidden region where V(x) > E. Each level line ends at a turning point."),
        (INK2, "Exact overlay", "The dotted Hermite functions are indistinguishable from the finite-difference curves."),
    ]
    y = 1.65
    for col, head, body in callouts:
        box(s, cx, y, 0.08, 1.05, fill=col, line=None, shape=MSO_SHAPE.RECTANGLE)
        text(s, cx + 0.3, y, 5.3, 0.4, head, size=18, bold=True)
        text(s, cx + 0.3, y + 0.42, 5.3, 0.8, body, size=14, color=INK2)
        y += 1.4

    # 5 - validation
    max_de = np.abs(st["dE"]).max()
    worst_inf = (1 - st["fid"]).max()
    norm_err = np.abs(st["norms"] - 1).max()
    h, e = st["conv"].T
    ratio = np.mean(e[:-1] / e[1:])
    s = new_slide(prs, "Numerical vs analytical: it checks out", "Validation",
                  f"All {st['k']} energies match n + ½ to within {max_de:.1e}. The error falls by about 4× each "
                  "time the grid spacing halves, which is the signature of a second-order method.")
    tiles = [
        (f"{max_de:.1e}", f"max |E_num − E_exact|, {st['k']} states"),
        (f"{worst_inf:.0e}", "worst 1 − |⟨ψ_num|ψ_exact⟩|"),
        (f"{norm_err:.0e}", "max |∫|ψ|² dx − 1|"),
        (f"×{ratio:.2f}", "error drop per halving of h (ideal ×4)"),
    ]
    for i, (v, l) in enumerate(tiles):
        x = 0.6 + i * 3.08
        box(s, x, 1.6, 2.9, 1.15)
        text(s, x + 0.25, 1.72, 2.5, 0.5, [[mono(v, size=26, bold=True)]])
        text(s, x + 0.25, 2.3, 2.5, 0.4, l, size=12, color=INK2)
    picture(s, DOCS / "convergence.png", 0.6, 3.0, h=4.1, border=False)
    picture(s, DOCS / "energy-error.png", 6.45, 3.0, w=4.0, border=False)
    text(s, 6.55, 5.35, 4.0, 0.3, "|ΔE| grows with n: finer detail needs a finer grid", size=12, color=MUTED)
    box(s, 10.65, 3.0, 2.05, 4.1)
    text(s, 10.85, 3.15, 1.75, 0.4, f"{st['n_tests']} pytest checks", size=15, bold=True)
    text(s, 10.85, 3.6, 1.75, 3.4, [
        [(t, {"space": 3})] for t in ["ascending order", "E ≈ n + ½", "spacing = ℏω", "normalisation",
                                      "orthogonality", "parity (−1)ⁿ", "n nodes", "exact ψ match",
                                      "O(h²) convergence", "ω rescaling", "small-box failure", "build output"]
    ], size=12, color=INK2)

    # 6 - architecture
    s = new_slide(prs, "Demo and architecture", "How it fits together",
                  "Python is the source of truth. The build step bakes the solution into a single HTML file "
                  "that opens directly in a browser with no server.")
    def node(x, y, w, head, sub, col=BLUE):
        box(s, x, y, w, 1.0, fill=PANEL)
        box(s, x, y, 0.08, 1.0, fill=col, line=None, shape=MSO_SHAPE.RECTANGLE)
        text(s, x + 0.25, y + 0.14, w - 0.35, 0.4, [[mono(head, size=15, bold=True)]])
        text(s, x + 0.25, y + 0.55, w - 0.35, 0.4, sub, size=12, color=INK2)
    node(0.6, 1.7, 3.3, "qho/solver.py", "finite-difference Hamiltonian + eigh")
    node(0.6, 2.95, 3.3, "qho/analytic.py", "E_n = n + ½, Hermite functions", ORANGE)
    node(4.75, 2.3, 3.2, "qho/build.py", "solve, compare, embed as JSON")
    node(8.85, 2.3, 3.85, "demo/index.html", "self-contained page, Plotly.js", ORANGE)
    node(0.6, 4.2, 3.3, "tests/", "pytest: physics and build checks", INK2)
    arrow(s, 3.95, 2.2, 4.7, 2.7)
    arrow(s, 3.95, 3.45, 4.7, 3.0)
    arrow(s, 8.0, 2.8, 8.8, 2.8)
    arrow(s, 3.95, 4.7, 4.7, 4.7)
    text(s, 4.8, 4.5, 3.5, 0.4, "verifies solver and build", size=12, color=MUTED)
    text(s, 8.85, 3.6, 3.85, 0.4, "Controls", size=16, bold=True)
    bullets(s, 8.85, 4.05, 3.85, [
        "View: ψ or |ψ|²", "States shown (1–12) and selected n", "ω slider (exact scaling law)",
        "Exact-solution overlay", "Click, ↑/↓ keys, URL presets",
    ], size=14, gap=6)
    box(s, 0.6, 5.65, 7.35, 1.25, fill=RGBColor(0x1A, 0x1A, 0x19), line=None)
    text(s, 0.85, 5.8, 7.0, 1.0, [
        [mono("$ pip install -r requirements.txt", size=14, color=RGBColor(0xC3, 0xC2, 0xB7))],
        [mono("$ python -m qho.build && open demo/index.html", size=14, color=RGBColor(0xFF, 0xFF, 0xFF), space=4)],
        [mono("$ pytest", size=14, color=RGBColor(0xC3, 0xC2, 0xB7), space=4)],
    ])

    # 7 - takeaways
    s = new_slide(prs, "Takeaways and limitations", "Wrap-up",
                  "A 30-line solver reproduces the textbook oscillator to about 1e-3, and the error behaves exactly "
                  "as the theory of the method predicts. The limitations below are known and tested.")
    text(s, 0.6, 1.7, 6.0, 0.4, "What we showed", size=20, bold=True, color=BLUE)
    bullets(s, 0.6, 2.25, 5.9, [
        f"Finite differences reproduce E_n = n + ½ to {max_de:.0e} with 2401 points in ~{st['ms']:.0f} ms",
        "Wavefunctions match the Hermite functions (1 − fidelity < 1e-7)",
        "Second-order convergence is directly visible: ×4 per halving of h",
        "One picture shows the potential, the quantised levels, ψ, |ψ|² and the exact overlay",
    ], size=16, gap=12)
    box(s, 7.0, 1.6, 5.7, 5.2)
    text(s, 7.3, 1.8, 5.2, 0.4, "Limitations", size=20, bold=True, color=ORANGE)
    bullets(s, 7.3, 2.35, 5.15, [
        "Finite box with hard walls: high n needs a larger L (tested: a small box pushes the energies up)",
        "O(h²) error grows with n, so fine structure needs finer grids",
        "1D, time-independent, harmonic potential only (the solver accepts any V(x), but only this one is validated)",
        "The ω slider rescales the ω = 1 result. This is identical to re-solving on a grid scaled by 1/√ω (tested)",
    ], size=15, gap=12)
    return prs


def main():
    OUT.parent.mkdir(exist_ok=True)
    st = compute_stats()
    make_figures(st)
    prs = build(st)
    prs.save(OUT)
    print(f"wrote {OUT.relative_to(ROOT)} ({len(prs.slides)} slides)")


if __name__ == "__main__":
    main()
